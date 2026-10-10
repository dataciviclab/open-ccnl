#!/usr/bin/env python3
"""Scarica l'Excel CNEL "dati sull'applicazione dei CCNL dal 2018" e ne estrae
il foglio quantitativo (datori + dipendenti per anno, con F/M dal 2023).

Uso: python3 extract.py raw.csv
"""

from __future__ import annotations

import csv
import io
import sys
import urllib.request
from pathlib import Path

EXCEL_URL = (
    "https://www.cnel.it/media/docs/default-source/"
    "archivio-contratti---cartelle-excel/"
    "dati-sull%27applicazione-dei-ccnl-dal-2018.xlsx"
)

YEARS = list(range(2018, 2026))

COLUMNS = [
    "codice_ccnl",
    "sottosezione",
    "status",
    "tipo_ccnl",
    "titolo",
    "firmatari_datoriali",
    "firmatari_sindacali",
    "anno",
    "n_datori",
    "n_dipendenti",
    "n_dipendenti_f",
    "n_dipendenti_m",
    "filtri_conto_annuale_ppaa",
    "fonte_note",
]


def _norm_header(value: object) -> str:
    if value is None:
        return ""
    # "n° datori di lavoro  2021" → "n° datori di lavoro 2021"
    return " ".join(str(value).lower().split())


def _to_int(value: object) -> int | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    text = str(value).strip().replace(".", "").replace(",", ".")
    if not text or text == "-":
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _to_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def download_xlsx() -> bytes:
    print(f"  download {EXCEL_URL}", flush=True)
    req = urllib.request.Request(EXCEL_URL, headers={"User-Agent": "dataciviclab-open-ccnl/0.1"})
    with urllib.request.urlopen(req, timeout=180) as resp:
        data = resp.read()
    print(f"  {len(data)} bytes")
    return data


def _resolve_year_columns(headers: list[str]) -> dict[int, dict[str, int]]:
    """Mappa anno → {datori, dipendenti, dipendenti_f, dipendenti_m} → col index."""
    out: dict[int, dict[str, int]] = {y: {} for y in YEARS}
    for j, h in enumerate(headers):
        if not h:
            continue
        for year in YEARS:
            ystr = str(year)
            if ystr not in h:
                continue
            if h.startswith("n° datori di lavoro"):
                out[year]["datori"] = j
            elif "dipendenti f" in h:
                out[year]["dipendenti_f"] = j
            elif "dipendenti m" in h:
                out[year]["dipendenti_m"] = j
            elif h.startswith("n° dipendenti"):
                # totale (senza F/M) — "n° dipendenti 2023"
                rest = h.replace("n° dipendenti", "").replace(ystr, "").strip()
                if rest == "":
                    out[year]["dipendenti"] = j
    return out


def extract_quantitativi(xlsx_bytes: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    candidates = [s for s in wb.sheetnames if "quantitativ" in s.lower()]
    if not candidates:
        raise RuntimeError(f"Foglio quantitativi non trovato. Sheets: {wb.sheetnames}")
    ws = wb[candidates[0]]

    rows_iter = ws.iter_rows(values_only=True)
    raw_header = next(rows_iter)
    headers = [_norm_header(h) for h in (raw_header or ())]
    idx = {h: i for i, h in enumerate(headers) if h}

    def find_exact(*options: str) -> int | None:
        for opt in options:
            if opt in idx:
                return idx[opt]
        return None

    i_code = find_exact("ccnl id")
    if i_code is None:
        raise RuntimeError(f"Colonna 'CCNL id' non trovata. Header: {headers[:12]}")

    i_sub = find_exact("settore privato e pubblico - sottosezione")
    i_status = find_exact("status")
    i_tipo = find_exact("tipo ccnl")
    i_titolo = find_exact("ccnl titolo")
    i_firm_d = find_exact("firmatari datoriali")
    i_firm_s = find_exact("firmatari sindacali")
    i_note = find_exact("filtri su conto annuale ppaa")
    i_fonte = find_exact("fonte dei dati e altre note")

    year_cols = _resolve_year_columns(headers)
    if not any(year_cols[y].get("dipendenti") for y in YEARS):
        raise RuntimeError(f"Nessuna colonna dipendenti per anno trovata. Header sample: {headers}")

    out: list[dict] = []
    for row in rows_iter:
        if row is None:
            continue
        code = _to_str(row[i_code] if i_code < len(row) else None)
        if not code:
            continue

        base = {
            "codice_ccnl": code,
            "sottosezione": _to_str(row[i_sub]) if i_sub is not None and i_sub < len(row) else "",
            "status": _to_str(row[i_status]) if i_status is not None and i_status < len(row) else "",
            "tipo_ccnl": _to_str(row[i_tipo]) if i_tipo is not None and i_tipo < len(row) else "",
            "titolo": _to_str(row[i_titolo]) if i_titolo is not None and i_titolo < len(row) else "",
            "firmatari_datoriali": _to_str(row[i_firm_d]) if i_firm_d is not None and i_firm_d < len(row) else "",
            "firmatari_sindacali": _to_str(row[i_firm_s]) if i_firm_s is not None and i_firm_s < len(row) else "",
            "filtri_conto_annuale_ppaa": _to_str(row[i_note]) if i_note is not None and i_note < len(row) else "",
            "fonte_note": _to_str(row[i_fonte]) if i_fonte is not None and i_fonte < len(row) else "",
        }

        for year in YEARS:
            cols = year_cols[year]
            if not cols:
                continue

            def cell(key: str) -> int | None:
                j = cols.get(key)
                if j is None or j >= len(row):
                    return None
                return _to_int(row[j])

            out.append(
                {
                    **base,
                    "anno": year,
                    "n_datori": cell("datori"),
                    "n_dipendenti": cell("dipendenti"),
                    "n_dipendenti_f": cell("dipendenti_f"),
                    "n_dipendenti_m": cell("dipendenti_m"),
                }
            )

    wb.close()
    return out


def main(out_path: Path) -> None:
    xlsx = download_xlsx()
    rows = extract_quantitativi(xlsx)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Totale: {len(rows)} righe -> {out_path}")


if __name__ == "__main__":
    dest = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("raw.csv")
    main(dest)
