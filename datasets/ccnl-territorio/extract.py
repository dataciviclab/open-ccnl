#!/usr/bin/env python3
"""Scarica l'Excel CNEL "dati applicazione CCNL" e estrae i fogli provinciali
INPS (provincia × genere) per 2023, 2024, 2025 in un unico CSV lungo.

Uso: python3 extract.py raw.csv
"""

from __future__ import annotations

import csv
import io
import sys
import urllib.request
from pathlib import Path

EXCEL_URL = (
    "https://www.cnel.it/media/docs/default-source/"
    "archivio-contratti---cartelle-excel/"
    "dati-sull%27applicazione-dei-ccnl-dal-2018.xlsx"
)

# foglio → anno (header variabile: CODICE_CCNL vs CODICE_CONTRATTO)
SHEET_BY_YEAR = {
    2023: "2-dati INPS 2023 provinc genere",
    2024: "3-dati INPS 2024 provinc genere",
    2025: "4 - dati INPS 2025 prov  genere",
}

COLUMNS = [
    "anno",
    "codice_ccnl",
    "regione",
    "sigla_provincia",
    "provincia",
    "sesso",
    "n_dipendenti",
]


def _norm(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _to_int(value: object) -> int | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def download_xlsx() -> bytes:
    print(f"  download {EXCEL_URL}", flush=True)
    req = urllib.request.Request(EXCEL_URL, headers={"User-Agent": "dataciviclab-open-ccnl/0.1"})
    with urllib.request.urlopen(req, timeout=180) as resp:
        data = resp.read()
    print(f"  {len(data)} bytes")
    return data


def extract_provinciali(xlsx_bytes: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    out: list[dict] = []

    for year, sheet_name in SHEET_BY_YEAR.items():
        # tolleranza su naming foglio
        if sheet_name not in wb.sheetnames:
            candidates = [
                s
                for s in wb.sheetnames
                if str(year) in s and ("prov" in s.lower() or "provinc" in s.lower())
            ]
            if not candidates:
                raise RuntimeError(
                    f"Foglio provinciale {year} non trovato. Sheets: {wb.sheetnames}"
                )
            sheet_name = candidates[0]

        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter)
        headers = [_norm(h).lower() for h in (header or ())]

        def find(*frags: str) -> int | None:
            for i, h in enumerate(headers):
                if all(f in h for f in frags):
                    return i
            return None

        i_code = find("codice")
        i_reg = find("regione")
        i_sig = find("sigla")
        i_prov = find("provincia")
        # provincia colonna è quella con nome lungo non sigla
        if i_prov is not None and i_sig is not None and i_prov == i_sig:
            for i, h in enumerate(headers):
                if "provincia" in h and "sigla" not in h:
                    i_prov = i
                    break
        i_sex = find("sesso")
        i_dip = find("dipendenti")
        if None in (i_code, i_reg, i_sig, i_prov, i_sex, i_dip):
            raise RuntimeError(f"Header provinciale {year} non riconosciuto: {headers}")

        n = 0
        for row in rows_iter:
            if row is None:
                continue
            code = _norm(row[i_code]) if i_code < len(row) else ""
            if not code:
                continue
            dip = _to_int(row[i_dip]) if i_dip < len(row) else None
            out.append(
                {
                    "anno": year,
                    "codice_ccnl": code.upper(),
                    "regione": _norm(row[i_reg]) if i_reg < len(row) else "",
                    "sigla_provincia": _norm(row[i_sig]).upper() if i_sig < len(row) else "",
                    "provincia": _norm(row[i_prov]) if i_prov < len(row) else "",
                    "sesso": _norm(row[i_sex]).upper() if i_sex < len(row) else "",
                    "n_dipendenti": dip if dip is not None else 0,
                }
            )
            n += 1
        print(f"  foglio {year}: {n} righe")

    wb.close()
    return out


def main(out_path: Path) -> None:
    rows = extract_provinciali(download_xlsx())
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Totale: {len(rows)} righe -> {out_path}")


if __name__ == "__main__":
    dest = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("raw.csv")
    main(dest)
